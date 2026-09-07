"""
Client-ready report rendering: DOCX and PDF.

Takes the enriched report context (the engagement's ROE, execution log, and
findings — with CVSS/severity) and produces a branded, structured document:
cover page, executive summary with a findings-by-severity table, findings
detail (description / impact / reproduction / remediation / evidence), and the
execution log. DOCX is the editable consultant deliverable; PDF is the
read-only counterpart.

Pure formatting of data already entered — nothing here runs or contacts a target.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone

_SEVERITY_ORDER = {"critical": 0, "high": 1, "medium": 2, "low": 3,
                   "informational": 4, None: 5}
# Cartographic palette — matches the app UI (survey terracotta / teal, map ink).
_ACCENT = "B0472C"     # survey terracotta
_TEAL = "3E6B72"       # survey teal
_INK = "20261E"        # topographic ink
_SEVERITY_COLOR = {
    "critical": "9C2B1B", "high": "C0562E", "medium": "B08814",
    "low": "3E7A5A", "informational": "3E6B72",
}
_SEV_LABEL = {"critical": "Critical", "high": "High", "medium": "Medium",
              "low": "Low", "informational": "Informational"}


def _sorted_findings(findings: list[dict]) -> list[dict]:
    return sorted(findings, key=lambda f: (_SEVERITY_ORDER.get(f.get("severity"), 5),
                                           str(f.get("title") or "")))


def _severity_counts(findings: list[dict]) -> dict[str, int]:
    counts = {k: 0 for k in ("critical", "high", "medium", "low", "informational")}
    for f in findings:
        sev = f.get("severity")
        if sev in counts:
            counts[sev] += 1
    return counts


def _now_str() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def _lbl(v) -> str:
    return str(v).replace("_", " ").replace("-", " ").title() if v else "—"


# ── DOCX ──────────────────────────────────────────────────────────

def to_docx(ctx: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    eng = ctx["engagement"]
    roe = ctx.get("roe", {})
    cov = ctx.get("coverage", {})
    findings = _sorted_findings(ctx.get("findings", []))
    counts = _severity_counts(findings)
    org_name = ctx.get("org_name") or "—"

    doc = Document()

    def _h1(text: str):
        h = doc.add_heading(level=1)
        r = h.add_run(text)
        r.font.color.rgb = RGBColor.from_string(_ACCENT)
        return h

    # Cover.
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(eng.get("name") or "Engagement Report")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor.from_string(_ACCENT)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Penetration Test Report").font.size = Pt(14)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"Client: {eng.get('client') or '—'}    ·    Prepared by: {org_name}\n"
        f"Generated {_now_str()}"
    ).font.size = Pt(10)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run("CONFIDENTIAL — Authorized assessment record. Planning and documentation only.")
    r.italic = True
    r.font.size = Pt(9)
    doc.add_page_break()

    # Overview.
    _h1("Overview")
    for k, v in (("Client", eng.get("client")), ("Authorization", eng.get("authorization_ref")),
                 ("Objective", _lbl(eng.get("objective"))), ("Box type", _lbl(eng.get("box_type"))),
                 ("Status", _lbl(eng.get("status")))):
        p = doc.add_paragraph()
        p.add_run(f"{k}: ").bold = True
        p.add_run(str(v) if v else "—")

    # Rules of engagement.
    _h1("Rules of engagement")
    for k, v in (("In-scope platforms", roe.get("scope_platforms")),
                 ("In-scope targets", roe.get("in_scope_targets")),
                 ("Restrictions", roe.get("restrictions")),
                 ("Time budget (h)", roe.get("time_budget_hours"))):
        p = doc.add_paragraph()
        p.add_run(f"{k}: ").bold = True
        p.add_run(", ".join(map(str, v)) if isinstance(v, list) else (str(v) if v else "—"))

    # Executive summary.
    _h1("Executive summary")
    doc.add_paragraph(
        f"Execution coverage: {cov.get('coverage_pct', 0)}% "
        f"({cov.get('steps_worked', 0)} of {cov.get('total_steps', 0)} planned steps worked). "
        f"{len(findings)} finding(s) recorded."
    )
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Light Grid Accent 1"
    hdr = tbl.rows[0].cells
    hdr[0].paragraphs[0].add_run("Severity").bold = True
    hdr[1].paragraphs[0].add_run("Count").bold = True
    for sev in ("critical", "high", "medium", "low", "informational"):
        row = tbl.add_row().cells
        run = row[0].paragraphs[0].add_run(_SEV_LABEL[sev])
        run.font.color.rgb = RGBColor.from_string(_SEVERITY_COLOR[sev])
        run.bold = True
        row[1].paragraphs[0].add_run(str(counts[sev]))

    # Findings detail.
    _h1("Findings")
    if not findings:
        doc.add_paragraph("No findings recorded.")
    for i, f in enumerate(findings, 1):
        h = doc.add_heading(level=2)
        sev = f.get("severity")
        tag = f"[{_SEV_LABEL.get(sev, '—')}] " if sev else ""
        run = h.add_run(f"{i}. {tag}{f.get('title') or 'Untitled'}")
        if sev in _SEVERITY_COLOR:
            run.font.color.rgb = RGBColor.from_string(_SEVERITY_COLOR[sev])
        cvss = f.get("cvss_score")
        if cvss is not None:
            p = doc.add_paragraph()
            p.add_run("CVSS: ").bold = True
            p.add_run(f"{cvss}" + (f"  ({f.get('cvss_vector')})" if f.get("cvss_vector") else ""))
        for label, key in (("Status", "status"), ("Affected assets", "affected_assets"),
                           ("Description", "description"), ("Impact", "impact"),
                           ("Reproduction", "reproduction"), ("Remediation", "remediation"),
                           ("ATT&CK techniques", "technique_ids"), ("References", "references")):
            val = f.get(key)
            if not val:
                continue
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(", ".join(map(str, val)) if isinstance(val, list) else str(val))
        ev = f.get("evidence") or []
        if ev:
            p = doc.add_paragraph()
            p.add_run("Evidence: ").bold = True
            p.add_run("; ".join(
                f"{e.get('filename','?')} (sha256 {str(e.get('sha256',''))[:12]}…)" for e in ev))

    # Execution log.
    _h1("Execution log")
    steps = ctx.get("steps", [])
    if not steps:
        doc.add_paragraph("No steps documented.")
    for s in steps:
        h = doc.add_heading(level=3)
        h.add_run(f"{s.get('index', 0) + 1}. {s.get('name') or s.get('technique_id')} — {_lbl(s.get('outcome'))}")
        meta_bits = [s.get("technique_id"), _lbl(s.get("tactic"))]
        if s.get("operator"):
            meta_bits.append(f"operator {s['operator']}")
        doc.add_paragraph(" · ".join(b for b in meta_bits if b)).italic = True
        if s.get("notes"):
            doc.add_paragraph(s["notes"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── PDF ────────────────────────────────────────────────────────────

def to_pdf(ctx: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )
    from xml.sax.saxutils import escape

    eng = ctx["engagement"]
    roe = ctx.get("roe", {})
    cov = ctx.get("coverage", {})
    findings = _sorted_findings(ctx.get("findings", []))
    counts = _severity_counts(findings)
    org_name = ctx.get("org_name") or "—"

    styles = getSampleStyleSheet()
    h1, h2, body = styles["Heading1"], styles["Heading2"], styles["BodyText"]
    center = ParagraphStyle("center", parent=body, alignment=TA_CENTER)
    cover_title = ParagraphStyle("cover", parent=styles["Title"], fontSize=26, spaceAfter=12,
                                textColor=colors.HexColor("#"+_ACCENT))
    h1.textColor = colors.HexColor("#"+_ACCENT)

    def esc(v) -> str:
        return escape(str(v)) if v is not None else "—"

    story: list = []
    story.append(Spacer(1, 1.6 * inch))
    story.append(Paragraph(esc(eng.get("name") or "Engagement Report"), cover_title))
    story.append(Paragraph("Penetration Test Report", center))
    story.append(Spacer(1, 0.3 * inch))
    story.append(Paragraph(
        f"Client: {esc(eng.get('client'))} &nbsp;·&nbsp; Prepared by: {esc(org_name)}", center))
    story.append(Paragraph(f"Generated {esc(_now_str())}", center))
    story.append(Spacer(1, 0.3 * inch))
    story.append(Paragraph(
        "<i>CONFIDENTIAL — Authorized assessment record. Planning and documentation only.</i>",
        center))
    story.append(PageBreak())

    story.append(Paragraph("Overview", h1))
    for k, v in (("Client", eng.get("client")), ("Authorization", eng.get("authorization_ref")),
                 ("Objective", _lbl(eng.get("objective"))), ("Box type", _lbl(eng.get("box_type"))),
                 ("Status", _lbl(eng.get("status")))):
        story.append(Paragraph(f"<b>{k}:</b> {esc(v)}", body))

    story.append(Paragraph("Rules of engagement", h1))
    for k, v in (("In-scope platforms", roe.get("scope_platforms")),
                 ("In-scope targets", roe.get("in_scope_targets")),
                 ("Restrictions", roe.get("restrictions")),
                 ("Time budget (h)", roe.get("time_budget_hours"))):
        val = ", ".join(map(str, v)) if isinstance(v, list) else v
        story.append(Paragraph(f"<b>{k}:</b> {esc(val)}", body))

    story.append(Paragraph("Executive summary", h1))
    story.append(Paragraph(
        f"Execution coverage: {cov.get('coverage_pct', 0)}% "
        f"({cov.get('steps_worked', 0)} of {cov.get('total_steps', 0)} planned steps worked). "
        f"{len(findings)} finding(s) recorded.", body))
    data = [["Severity", "Count"]] + [
        [_SEV_LABEL[s], str(counts[s])] for s in
        ("critical", "high", "medium", "low", "informational")]
    tbl = Table(data, colWidths=[2.5 * inch, 1.0 * inch])
    tstyle = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#"+_TEAL)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]
    for i, s in enumerate(("critical", "high", "medium", "low", "informational"), start=1):
        tstyle.append(("TEXTCOLOR", (0, i), (0, i), colors.HexColor("#" + _SEVERITY_COLOR[s])))
    tbl.setStyle(TableStyle(tstyle))
    story.append(Spacer(1, 6))
    story.append(tbl)

    story.append(Paragraph("Findings", h1))
    if not findings:
        story.append(Paragraph("No findings recorded.", body))
    for i, f in enumerate(findings, 1):
        sev = f.get("severity")
        color = "#" + _SEVERITY_COLOR.get(sev, "222222")
        tag = f"[{_SEV_LABEL.get(sev, '—')}] " if sev else ""
        story.append(Paragraph(
            f'<font color="{color}"><b>{i}. {esc(tag)}{esc(f.get("title") or "Untitled")}</b></font>', h2))
        if f.get("cvss_score") is not None:
            story.append(Paragraph(
                f"<b>CVSS:</b> {esc(f.get('cvss_score'))} {esc(f.get('cvss_vector') or '')}", body))
        for label, key in (("Status", "status"), ("Affected assets", "affected_assets"),
                           ("Description", "description"), ("Impact", "impact"),
                           ("Reproduction", "reproduction"), ("Remediation", "remediation"),
                           ("ATT&CK techniques", "technique_ids"), ("References", "references")):
            val = f.get(key)
            if not val:
                continue
            shown = ", ".join(map(str, val)) if isinstance(val, list) else val
            story.append(Paragraph(f"<b>{label}:</b> {esc(shown)}", body))
        ev = f.get("evidence") or []
        if ev:
            shown = "; ".join(
                f"{e.get('filename','?')} (sha256 {str(e.get('sha256',''))[:12]}…)" for e in ev)
            story.append(Paragraph(f"<b>Evidence:</b> {esc(shown)}", body))
        story.append(Spacer(1, 6))

    story.append(Paragraph("Execution log", h1))
    steps = ctx.get("steps", [])
    if not steps:
        story.append(Paragraph("No steps documented.", body))
    for s in steps:
        story.append(Paragraph(
            f"<b>{s.get('index', 0) + 1}. {esc(s.get('name') or s.get('technique_id'))} — "
            f"{_lbl(s.get('outcome'))}</b>", h2))
        meta_bits = [s.get("technique_id"), _lbl(s.get("tactic"))]
        if s.get("operator"):
            meta_bits.append(f"operator {s['operator']}")
        story.append(Paragraph("<i>" + esc(" · ".join(b for b in meta_bits if b)) + "</i>", body))
        if s.get("notes"):
            story.append(Paragraph(esc(s["notes"]), body))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=letter,
                      leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                      topMargin=0.9 * inch, bottomMargin=0.8 * inch).build(story)
    return buf.getvalue()
